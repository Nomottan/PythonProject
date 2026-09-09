from pathlib import Path
import openpyxl
import xlrd
import csv
import re
from openpyxl import Workbook
from typing import List, Dict, Tuple
from utils.text_utils import TextUtils

class ExcelHelper:
    @staticmethod
    def get_column_index(sheet, header_name: str) -> int | None:
        if sheet.max_row < 1:
            return None
        header_row = sheet[1]
        for idx, cell in enumerate(header_row, start=1):
            if cell.value and str(cell.value).strip().lower() == header_name.strip().lower():
                return idx
        return None

    @staticmethod
    def open_data_file(file_path, read_only=True, data_only=True):
        """
        Открывает файл любого поддерживаемого формата (.xlsx, .xls, .csv) и возвращает объект,
        который имеет интерфейс, аналогичный openpyxl Workbook (sheetnames, active, iter_rows).
        """
        file_path = Path(file_path)
        ext = file_path.suffix.lower()

        if ext in ('.xlsx', '.xlsm'):
            return openpyxl.load_workbook(file_path, read_only=read_only, data_only=data_only)
        elif ext == '.xls':
            wb = xlrd.open_workbook(file_path, formatting_info=False, on_demand=True)
            return XlsReader(wb)
        elif ext == '.csv':
            return CsvReader(file_path)
        else:
            raise ValueError(f"Неподдерживаемый формат файла: {ext}")

    @staticmethod
    def get_first_row(sheet):
        if sheet.max_row < 1:
            return []
        return [cell.value for cell in sheet[1]]

    @staticmethod
    def find_sheet_by_key(wb, keys):
        sheets_normalized = [TextUtils.normalize(s) for s in wb.sheetnames]
        for key in keys:
            norm_key = TextUtils.normalize(key)
            if norm_key in sheets_normalized:
                idx = sheets_normalized.index(norm_key)
                return wb.sheetnames[idx]
        return None

    @staticmethod
    def open_workbook_safe(file_path, read_only=False, data_only=True):
        try:
            from openpyxl import load_workbook
            return load_workbook(file_path, data_only=data_only, read_only=read_only)
        except Exception:
            return None

    @staticmethod
    def create_workbook_with_headers(headers, sheet_name="Лист1", write_only=False):
        wb = Workbook(write_only=write_only)
        if write_only:
            ws = wb.create_sheet(sheet_name)
        else:
            ws = wb.active
            ws.title = sheet_name
        if headers:
            ws.append(headers)
        return wb, ws

    @staticmethod
    def append_headers(worksheet, headers):
        worksheet.append(headers)

    @staticmethod
    def copy_filtered_rows(source_sheet, target_sheet, columns_to_keep,
                           condition, start_row=2, status_col_idx=None, status_counts=None):
        rows_copied = 0
        for row in source_sheet.iter_rows(min_row=start_row, values_only=True):
            if not condition(row):
                continue
            new_row = []
            for col_idx in columns_to_keep:
                value = row[col_idx - 1] if len(row) >= col_idx else None
                new_row.append(value)
            target_sheet.append(new_row)

            if status_col_idx is not None and status_counts is not None:
                status_val = row[status_col_idx - 1] if len(row) >= status_col_idx else None
                status_str = str(status_val).strip() if status_val is not None else ""
                if status_str:
                    status_counts[status_str] = status_counts.get(status_str, 0) + 1
                else:
                    status_counts["(пусто)"] = status_counts.get("(пусто)", 0) + 1

            rows_copied += 1
        return rows_copied

    @staticmethod
    def close_workbooks(workbooks_dict):
        if not workbooks_dict:
            return
        for value in workbooks_dict.values():
            if isinstance(value, list):
                for wb in value:
                    wb.close()
            else:
                value.close()
        workbooks_dict.clear()

    @staticmethod
    def open_workbook_with_ctx(file_path, ctx, description="файл", read_only=False, data_only=True):
        wb = ExcelHelper.open_workbook_safe(file_path, read_only=read_only, data_only=data_only)
        if wb is None:
            ctx.log(f"Ошибка открытия {description}: {Path(file_path).name}")
        return wb

    @staticmethod
    def find_seller_by_sheet_name(wb, sellers, sheet_name):
        normalized_sheet = TextUtils.normalize(sheet_name)
        for seller in sellers:
            if any(TextUtils.normalize(key) == normalized_sheet for key in seller.keys):
                return seller
        return None

    @staticmethod
    def parse_seller_from_mp_filename(filename):
        import re
        match = re.search(r"ОТЧЁТ МП ПО (.+?)\.xlsx$", filename, re.IGNORECASE)
        if match:
            candidate = match.group(1).strip()
            candidate = re.sub(r'\s+\d{1,2}[._-]\d{1,2}[._-]\d{2,4}$', '', candidate)
            candidate = candidate.strip()
            return candidate
        return None

    @staticmethod
    def read_column_values(sheet, col_index, start_row=2):
        values = []
        for row in sheet.iter_rows(min_row=start_row, values_only=True):
            if len(row) >= col_index and row[col_index - 1] is not None:
                val = str(row[col_index - 1]).strip()
                if val:
                    values.append(val)
        return values

    @staticmethod
    def filter_and_replace_sheet(file_path, sheet_name, conditions, ctx, description="файл"):
        wb = ExcelHelper.open_workbook_with_ctx(file_path, ctx, description=description, read_only=False,
                                                data_only=True)
        if wb is None:
            return None

        try:
            sheet = wb[sheet_name] if sheet_name else wb.active
            headers = [cell.value for cell in sheet[1]]
            rows = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                rows.append(row)

            total = len(rows)
            kept = []
            removed_stats = {label: 0 for label, _ in conditions}

            for row in rows:
                keep = True
                for label, func in conditions:
                    if not func(row):
                        removed_stats[label] += 1
                        keep = False
                        break
                if keep:
                    kept.append(row)

            new_wb, new_ws = ExcelHelper.create_workbook_with_headers(headers, sheet_name=sheet.title)
            for row in kept:
                new_ws.append(row)
            new_wb.save(file_path)
            new_wb.close()

            return {
                "total": total,
                "kept": len(kept),
                "removed": removed_stats
            }

        finally:
            wb.close()

    @staticmethod
    def filter_and_clean_rows(sheet, status_col=None, owner_col=None,
                              allowed_owners=None, required_status=None):
        filtered_rows = []
        stats = {'status': {}, 'owner': {}, 'total_removed': 0, 'total_kept': 0}

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            keep = True

            if status_col is not None and required_status is not None:
                status_val = row[status_col - 1] if len(row) >= status_col else None
                status_str = str(status_val).strip() if status_val is not None else ""
                if status_str.upper() != required_status.upper():
                    stats['status'][status_str or "(пусто)"] = stats['status'].get(status_str or "(пусто)", 0) + 1
                    keep = False

            if keep and owner_col is not None and allowed_owners is not None:
                owner_val = row[owner_col - 1] if len(row) >= owner_col else None
                owner_str = str(owner_val).strip() if owner_val is not None else ""
                if not owner_str:
                    stats['owner']["(пусто)"] = stats['owner'].get("(пусто)", 0) + 1
                    keep = False
                else:
                    norm_owner = TextUtils.normalize(owner_str)
                    if norm_owner not in allowed_owners:
                        stats['owner'][owner_str] = stats['owner'].get(owner_str, 0) + 1
                        keep = False

            if keep:
                filtered_rows.append((row_idx, row))
                stats['total_kept'] += 1
            else:
                stats['total_removed'] += 1

        return filtered_rows, stats

    @staticmethod
    def append_row_to_file(file_path, row_data, headers=None, sheet_name=None):
        import openpyxl
        file_path = Path(file_path)
        if not file_path.exists():
            if headers is None:
                return
            wb, ws = ExcelHelper.create_workbook_with_headers(headers, sheet_name=sheet_name or "Лист1")
            wb.save(file_path)
            wb.close()

        wb = openpyxl.load_workbook(file_path)
        ws = wb[sheet_name] if sheet_name else wb.active
        ws.append(row_data)
        wb.save(file_path)
        wb.close()

    @staticmethod
    def is_file_empty(file_path, sheet_name=None):
        try:
            import openpyxl
            file_path = Path(file_path)
            if not file_path.exists():
                return True
            wb = openpyxl.load_workbook(file_path)
            ws = wb[sheet_name] if sheet_name else wb.active
            is_empty = ws.max_row == 1
            wb.close()
            return is_empty
        except Exception:
            return True

class CsvReader:
    """Адаптер для чтения CSV-файлов как Excel-листа."""

    def __init__(self, file_path):
        self.file_path = Path(file_path)
        self._sheet_names = ['Sheet1']
        self._active_sheet = 'Sheet1'
        self._dialect = None
        self._max_row = None
        self._max_column = None
        self._data = None
        self._is_closed = False

    def _detect_dialect(self):
        if self._dialect is not None:
            return self._dialect
        with open(self.file_path, 'r', encoding='utf-8-sig') as f:
            sample = f.read(1024)
            f.seek(0)
            try:
                self._dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                self._dialect = csv.excel()
            return self._dialect

    def _load_data(self):
        if self._data is not None:
            return
        if self._is_closed:
            return
        dialect = self._detect_dialect()
        self._data = []
        with open(self.file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f, dialect)
            for row in reader:
                self._data.append(row)
        self._max_row = len(self._data)
        self._max_column = max((len(row) for row in self._data), default=0)

    def close(self):
        self._is_closed = True
        self._data = None

    @property
    def sheetnames(self):
        return self._sheet_names

    @property
    def active(self):
        return self

    @property
    def max_row(self):
        self._load_data()
        return self._max_row

    @property
    def max_column(self):
        self._load_data()
        return self._max_column

    def iter_rows(self, min_row=1, max_row=None, values_only=True):
        if not values_only:
            raise NotImplementedError("Только values_only=True поддерживается для CSV")
        self._load_data()
        if self._is_closed or self._data is None:
            return
        start = min_row - 1
        end = self._max_row if max_row is None else min(max_row, self._max_row)
        for row_idx in range(start, end):
            if row_idx < len(self._data):
                yield self._data[row_idx]

class XlsReader:
    """Адаптер для чтения .xls файлов через xlrd."""

    def __init__(self, workbook):
        self.workbook = workbook
        self._sheet_names = workbook.sheet_names()
        self._active = None
        self._is_closed = False

    def close(self):
        self._is_closed = True
        if hasattr(self.workbook, 'release_resources'):
            self.workbook.release_resources()

    @property
    def sheetnames(self):
        return self._sheet_names

    @property
    def active(self):
        if self._active is None and not self._is_closed:
            self._active = self.workbook.sheet_by_index(0)
        return self._active

    @property
    def max_row(self):
        if self._is_closed:
            return 0
        return self.active.nrows

    @property
    def max_column(self):
        if self._is_closed:
            return 0
        return self.active.ncols

    def iter_rows(self, min_row=1, max_row=None, values_only=True):
        if self._is_closed:
            return
        sheet = self.active
        if not values_only:
            raise NotImplementedError("Только values_only=True поддерживается для .xls")
        start = min_row - 1
        end = min(sheet.nrows, max_row) if max_row is not None else sheet.nrows
        for row_idx in range(start, end):
            yield sheet.row_values(row_idx)

class CsvNormalizer:
    """
    Нормализует CSV-файлы поставок: извлекает GTIN, КИЗ, название из всех ячеек строки.
    Валидирует данные и логирует некорректные строки.
    """

    def __init__(self, file_path: Path):
        self.file_path = Path(file_path)
        self.dialect = None

    def _detect_dialect(self):
        if self.dialect is not None:
            return self.dialect
        with open(self.file_path, 'r', encoding='utf-8-sig') as f:
            sample = f.read(1024)
            f.seek(0)
            try:
                self.dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                self.dialect = csv.excel()
            return self.dialect

    def normalize_rows(self, log_callback=None) -> Tuple[List[Dict[str, str]], List[str]]:
        """
        Нормализует данные по строкам: извлекает GTIN, КИЗ, название из всех ячеек строки.
        Возвращает (валидные_записи, список_некорректных_строк).
        """
        dialect = self._detect_dialect()
        rows = []
        with open(self.file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f, dialect)
            for row in reader:
                rows.append(row)

        if not rows:
            return [], []

        # Пропускаем заголовок (первая строка)
        data_rows = rows[1:] if len(rows) > 1 else []

        valid_rows = []
        invalid_log = []

        for row_idx, row in enumerate(data_rows, start=2):
            gtin = None
            kiz = None
            name = None
            count = None

            # Сканируем все ячейки строки
            for cell in row:
                if not cell:
                    continue
                cell_str = str(cell).strip()
                if not cell_str:
                    continue

                # Проверяем на GTIN (число 13-14 цифр, часто начинается с 46)
                if re.match(r'^46\d{11,12}$', cell_str):
                    gtin = cell_str
                    continue
                elif re.match(r'^\d{13,14}$', cell_str) and not gtin:
                    gtin = cell_str
                    continue

                # Проверяем на КИЗ (начинается с 01, длина > 31)
                if cell_str.startswith('01') and len(cell_str) > 31:
                    # Берём только до первого разделителя  (ASCII 29)
                    if '' in cell_str:
                        kiz = cell_str.split('')[0]
                    else:
                        kiz = cell_str
                    continue

                # Проверяем на название (длинный текст > 31 символов с пробелами)
                if len(cell_str) > 31 and ' ' in cell_str and not cell_str.isdigit():
                    name = cell_str
                    continue

                # Проверяем на количество (если есть отдельный числовой столбец)
                if cell_str.isdigit() and len(cell_str) <= 6 and not gtin:
                    # Если число меньше 10000 и это не GTIN, считаем количеством
                    if int(cell_str) < 10000:
                        count = cell_str

            # Валидация: строка валидна, если есть GTIN и хотя бы один из (КИЗ, название)
            if gtin and (kiz or name):
                record = {
                    'gtin': gtin,
                    'kiz': kiz if kiz else '',
                    'name': name if name else '',
                    'count': '1',  # по умолчанию 1
                    'source_file': self.file_path.name
                }
                valid_rows.append(record)
            else:
                # Формируем лог-сообщение
                missing = []
                if not gtin:
                    missing.append("GTIN")
                if not kiz and not name:
                    missing.append("КИЗ/название")
                invalid_log.append(f"Строка {row_idx}: отсутствуют {', '.join(missing)}")
                invalid_log.append(f"  Данные: {row[:10]}")  # только первые 10 ячеек для краткости

        return valid_rows, invalid_log