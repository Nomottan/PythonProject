import shutil
import openpyxl
from pathlib import Path
from datetime import date
from utils.context import TaskContext

class SalesAccumulatorService:
    """
    Сервис аккумуляции файлов продаж из двух папок в одну общую.
    Работает с новым форматом файлов продаж:
      - Имя: "{от_кого} - {кому} : {ИНН}.xlsx" (содержит " - " и " : ")
      - Столбцы: Наименование продукта, КИЗ (31 символов), GTIN, Цена.
    Приоритет отдаётся файлам из ЧЗ_МП (Sells_FBS):
      - если КИЗ присутствует в ЧЗ_МП, то строки из Возвратов с этим КИЗом игнорируются.
    """

    def accumulate(self, target_dir: str, first_folder: str, log_callback=None):
        """
        Аккумулирует файлы продаж из папок ЧЗ_МП и Возвраты в папку Продажи.
        Параметр first_folder игнорируется (всегда сначала ЧЗ_МП, затем Возвраты).
        """
        today = date.today()
        date_str = f"{today.day}_{today.month}_{today.year}"
        base_path = Path(target_dir) / date_str

        chz_folder = base_path / f"ЧЗ_МП_{date_str}"
        returns_folder = base_path / f"Возвраты_{date_str}"
        sales_folder = base_path / f"Продажи_{date_str}"

        ctx = TaskContext(target_dir, "Продажи_{date}", "log_аккумуляция.txt", log_callback)
        sales_folder.mkdir(parents=True, exist_ok=True)

        ctx.log("=== АККУМУЛЯЦИЯ ПРОДАЖ (приоритет ЧЗ_МП) ===")
        ctx.log(f"Рабочая папка: {sales_folder}")

        # 1. Обработка ЧЗ_МП (первая, приоритетная)
        chz_files = self._find_sales_files(chz_folder)
        kiz_from_fbs = set()  # множество КИЗов из ЧЗ_МП

        ctx.log(f"\n--- ОБРАБОТКА ЧЗ_МП ---")
        for src_path in chz_files:
            dst_path = sales_folder / src_path.name
            shutil.copy2(src_path, dst_path)
            kiz_set = self._collect_kiz_set(dst_path)
            kiz_from_fbs.update(kiz_set)
            ctx.log(f"  Скопирован: {src_path.name} (КИЗов: {len(kiz_set)})")

        # 2. Обработка Возвратов (с фильтрацией и сбором деталей)
        returns_files = self._find_sales_files(returns_folder)
        ctx.log(f"\n--- ОБРАБОТКА ВОЗВРАТОВ (фильтрация по КИЗам из ЧЗ_МП) ---")

        details = {
            "kiz_from_fbs": kiz_from_fbs,
            "files": []
        }

        for src_path in returns_files:
            dst_path = sales_folder / src_path.name
            src_kiz_set = self._collect_kiz_set(src_path)
            duplicates_in_file = src_kiz_set & kiz_from_fbs
            filtered_kiz = src_kiz_set - kiz_from_fbs

            file_info = {
                "name": src_path.name,
                "total": len(src_kiz_set),
                "duplicates": duplicates_in_file,
                "filtered": filtered_kiz
            }
            details["files"].append(file_info)

            if not filtered_kiz:
                ctx.log(
                    f"  Пропущен {src_path.name}: все КИЗы уже есть в ЧЗ_МП (всего {len(src_kiz_set)}, дубликатов {len(duplicates_in_file)})")
                continue

            # Дозапись или создание нового файла
            if dst_path.exists():
                # Дописываем строки в существующий файл
                wb_dst = openpyxl.load_workbook(dst_path)
                sheet_dst = wb_dst.active
                wb_src = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
                sheet_src = wb_src.active
                rows_added = 0
                for row in sheet_src.iter_rows(min_row=2, values_only=True):
                    # КИЗ теперь во втором столбце (индекс 1)
                    if len(row) >= 2 and row[1]:
                        kiz = str(row[1]).strip()
                        if kiz in filtered_kiz:
                            sheet_dst.append(row)
                            rows_added += 1
                wb_dst.save(dst_path)
                wb_dst.close()
                wb_src.close()
                ctx.log(
                    f"  Дополнен {dst_path.name}: добавлено {rows_added} строк (всего КИЗов {len(src_kiz_set)}, из них добавлено {len(filtered_kiz)})")
            else:
                # Создаём новый файл с заголовками и только отфильтрованными строками
                wb_dst = openpyxl.Workbook()
                sheet_dst = wb_dst.active
                wb_src = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
                sheet_src = wb_src.active
                header = list(sheet_src.iter_rows(min_row=1, max_row=1, values_only=True))[0]
                sheet_dst.append(header)
                rows_added = 0
                for row in sheet_src.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 2 and row[1]:
                        kiz = str(row[1]).strip()
                        if kiz in filtered_kiz:
                            sheet_dst.append(row)
                            rows_added += 1
                wb_dst.save(dst_path)
                wb_dst.close()
                wb_src.close()
                ctx.log(
                    f"  Создан новый файл {dst_path.name}: добавлено {rows_added} строк (всего КИЗов {len(src_kiz_set)}, из них добавлено {len(filtered_kiz)})")

        # Сохраняем детальный лог
        if details["files"]:
            log_path = sales_folder / "log_фильтрация_КИЗов.txt"
            self._log_kiz_details(log_path, details, ctx)

        ctx.log("\n=== АККУМУЛЯЦИЯ ЗАВЕРШЕНА ===")

    # ---------- Вспомогательные методы ----------

    def _find_sales_files(self, folder: Path) -> list[Path]:
        """
        Возвращает все файлы продаж в папке (новый формат).
        Отличительный признак: имя содержит " - " и " : ".
        """
        if not folder.exists():
            return []
        return [f for f in folder.glob("*.xlsx") if " - " in f.stem and " : " in f.stem]

    def _collect_kiz_set(self, file_path: Path) -> set:
        """
        Читает КИЗы из второго столбца (индекс 1) и возвращает множество.
        Структура нового файла: [Наименование, КИЗ, GTIN, Цена].
        """
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = wb.active
            kiz_set = set()
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if len(row) >= 2 and row[1]:  # второй столбец
                    kiz = str(row[1]).strip()
                    if kiz:
                        kiz_set.add(kiz)
            wb.close()
            return kiz_set
        except Exception:
            return set()

    def _log_kiz_details(self, log_path: Path, details: dict, ctx):
        """Сохраняет детальную информацию о фильтрации КИЗов для каждого файла возвратов."""
        with open(log_path, "w", encoding="utf-8") as f:
            f.write("=== ДЕТАЛИ ФИЛЬТРАЦИИ КИЗОВ ===\n\n")

            f.write("КИЗы из ЧЗ_МП (приоритетные):\n")
            if details["kiz_from_fbs"]:
                for kiz in sorted(details["kiz_from_fbs"]):
                    f.write(f"  {kiz}\n")
            else:
                f.write("  (нет)\n")
            f.write("\n" + "=" * 60 + "\n\n")

            for file_info in details["files"]:
                f.write(f"Файл: {file_info['name']}\n")
                f.write(f"  Всего КИЗов в файле: {file_info['total']}\n")
                f.write(f"  Дубликаты (уже есть в ЧЗ_МП): {len(file_info['duplicates'])}\n")
                if file_info['duplicates']:
                    for kiz in sorted(file_info['duplicates']):
                        f.write(f"    {kiz}\n")
                else:
                    f.write("    (нет)\n")
                f.write(f"  Отфильтрованные (добавлены): {len(file_info['filtered'])}\n")
                if file_info['filtered']:
                    for kiz in sorted(file_info['filtered']):
                        f.write(f"    {kiz}\n")
                else:
                    f.write("    (нет)\n")
                f.write("\n" + "-" * 40 + "\n")

        ctx.log(f"  Детальный лог фильтрации КИЗов сохранён в {log_path.name}")